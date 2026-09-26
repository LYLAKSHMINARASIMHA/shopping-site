from flask import Flask, render_template, request, url_for , redirect , session, jsonify
import random
import pandas as pd
import os
from openpyxl import load_workbook
from datetime import datetime
app = Flask(__name__)

@app.context_processor
def inject_user():
     return{
          "userid": session.get("userid"),
          "u_Name": session.get("u_Name"),
          "status_D": session.get("status_D"),
          "roll": session.get("roll")
     }

app.secret_key = "my_secret_key"

folders ={
    "mens": "static/products/mens_dresses", #size
    "womens": "static/products/women_dresses",#size
    "phones": "static/products/phones",  #No
    "slippers" : "static/products/slippers",  # size
    "shoes" : "static/products/shoes",  # size
    "watchs" : "static/products/watchs",  #size
    "products": "static/products",
}
df = pd.read_excel("Shopping_W_data.xlsx")
df.columns = df.columns.str.strip().str.lower()
df = df.dropna(how="all")
product = df.to_dict(orient="records")








# ............. product data 
product_dict ={}
for item in product:
     product_dict[str(item["p_id"])] = item


# ..........images path len(UHproduct)
imgpaths = {}
for root, dirs, files in os.walk("static/products"):
    for file in files:
        img_ID = os.path.splitext(file)[0]
        imgpaths[img_ID] = os.path.join(root,file).replace("\\","/").replace("static/","")


@app.route("/")
def Home():
    #............ order status update 
    if session.get("roll"):
        print(session.get("roll"))
    Pdate = datetime.now().date()
    
    OSdf = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_history")
    OSdf.columns = OSdf.columns.str.strip().str.lower()
    OSdf = OSdf.dropna(how="all")
    OSproduct = OSdf.to_dict(orient="records")


    if session.get("status_D") != str(Pdate):
        Orderstatus = {}
        for item in OSproduct:
            delivery_date = datetime.strptime( item["date"].split("||")[1], "%d/%m/%y").date()
            if item["status"] in ["Cancelled", "Delivered"]:
                continue
            else:
                if Pdate < delivery_date:
                    Orderstatus[item["orderid"]] = {
                        "Order status": "shipping"
                    }
                else:
                    Orderstatus[item["orderid"]] = {
                        "Order status": "Delivered"
                    }

        for orderid, data in Orderstatus.items():
    
            if orderid in OSdf["orderid"].values:
                OSdf.loc[OSdf["orderid"] == orderid, "status" ]= data["Order status"]
                print(f"{orderid} updated")
        with pd.ExcelWriter(
        "Shopping_W_data.xlsx",
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace"
        ) as writer:

            OSdf.to_excel(
                writer,
                sheet_name="users_history",
                index=False
            )

        print("Order status updated successfully.")

        session["status_D"] = str(Pdate)
    
    offmenu = True
    

    images = []
    for file in os.listdir(folders["mens"]):
        images.append("products/mens_dresses/" + file)
        images = images[:6]
    grouped_Bimg =[]
    for i in range(0, len(images),3):
        grouped_Bimg.append(images[i:i+3])
   
        # --------------------------------------------------

    Gimages = []
    for Gfile in os.listdir(folders["womens"]):
        Gimages.append("products/women_dresses/" + Gfile)
        Gimages = Gimages[:6]
    grouped_Gimg = []
    for i in range(0, len(Gimages),3):
            grouped_Gimg.append(Gimages[i:i+3])
        # ---------------------------------------
   
    Mimage = []
    for Mfile in os.listdir(folders["phones"]):
        Mimage.append("/products/phones/" + Mfile)
        Mimage = Mimage[:6]
    grouped_Mimg = []
    for i in range(0, len(Mimage), 3):
            grouped_Mimg.append(Mimage[i:i+3])
   

    return render_template("Home.html",
                            grouped_Bimg=grouped_Bimg,
                              grouped_Gimg = grouped_Gimg,
                                grouped_Mimg = grouped_Mimg,
                                offmenu = offmenu,
                                )

@app.route("/Orders/")
def Orders():

    #  users_history sheet
    OdataF = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_history")
    OdataF.columns = OdataF.columns.str.strip().str.lower()
    OdataF = OdataF.dropna(how="all")
    UHproduct = OdataF.to_dict(orient="records")

    # ..........images path len(UHproduct)
    imgpaths = {}
    for root, dirs, files in os.walk("static/products"):
        for file in files:
            img_ID = os.path.splitext(file)[0]
            imgpaths[img_ID] = os.path.join(root,file).replace("\\","/").replace("static/","")
    
    Pdate = datetime.now()
    Pdate = Pdate.strftime("%d/%m/%y")
    userID =session.get("userid")
    p_data = []
    OrderID =[] 
    Pimgid=[]
    UOdate=[]

    
    for item in UHproduct:
        if str(userID) == str(item["userid"]):
            OrderID.append({
                 "OrderID" : item["orderid"]
            })
            p_data.append({
                      "quantity": item["quantity"],
                      "status": item["status"]
            })
            UOdate.append({
                 "order date":item["date"].split("||")[0],
                 "delivery date":item["date"].split("||")[1]
            })
            Pimgid.append(item["p_id"])
            
    O_ID =[]
    for item in UHproduct:
        if str(userID) == str(item["userid"]):
            O_ID.append({
                 "OrderID" : item["orderid"]
            })

    

    

    imgdata = {}
    for pid in Pimgid:
        if str(pid) in product_dict:
            imgdata[pid] = product_dict[pid]

    imglocation = {}
    for item in Pimgid:
         imglocation[item] = imgpaths[item]
         
    Orderdata=[]
    for i, pid in enumerate(Pimgid):
         Orderdata.append({
              "OrderID":O_ID[i]["OrderID"],
              "image":imglocation[pid],
              "product":imgdata[pid],
              "order_date":UOdate[i]["order date"],
              "delivery_date":UOdate[i]["delivery date"],
              "totalamount":p_data[i]["quantity"]*imgdata[pid]["p_price"],
              "status":p_data[i]["status"],
              "quantity":p_data[i]["quantity"]
         })

         

    return render_template("Orders.html",
                              UOdate = UOdate,
                              Orderdata = Orderdata
                            )



@app.route("/Help_Center")
def Help_Center():
    return render_template("Help_Center.html")

@app.route("/searchproducts")
def searchproducts():

    search_words = {
    "men": "mens",
    "mens": "mens",
    "men wear": "mens",
    "mens wear": "mens",
    "women": "womens",
    "womens": "womens",
    "women wear": "womens",
    "womens wear": "womens",
    "women dress": "womens",
    "womens dress": "womens",
    "mobile": "phones",
    "mobiles": "phones",
    "watch": "watchs",
    "watches": "watchs",
    "mens watchs": "watchs",
    "womens watchs": "watchs",
    "men watchs": "watchs",
    "women watchs": "watchs",
    "sports shoes": "shoes",
    "running shoes": "shoes",
    }
    
    data = request.args.get("search").lower()
    data = data.replace("'","")
    data = search_words.get(data, data)
    
    matchedID =[]
    matchedID2 =[]
    searchimg = []
    for item in folders:
        if data == item:
            matchedID.append(item)


    if matchedID:
         for item in matchedID:
                 folder_path = folders[item].replace("static/","")
         
         for item in matchedID:
           for file in os.listdir(folders[item]):
                    matchedID2.append(folder_path+"/"+ file)

         selectSHimg = (random.sample(matchedID2, 6))
         for i in range(0, len(selectSHimg), 3):
             searchimg.append(selectSHimg[i:i+3])
    elif len(matchedID) >= 6:
        for item1 in product:
            if data in item1["category"]:
                matchedID.append(item1["p_id"])
        if len(matchedID) >= 6:
                print(matchedID)
        else:
             for item2 in product:
                if data in item2["p_description"]:
                    matchedID.append(item2["p_id"])
             print(matchedID)
             
        

    

    

    


    offmenu = True
    imagename="products/img/"

    imageId = []
    for file in os.listdir(folders["mens"]):
        imageId.append("products/mens_dresses/" + file)
    
    for file in os.listdir(folders["womens"]):
        imageId.append("products/women_dresses/" + file)

    for file in os.listdir(folders["phones"]):
        imageId.append("products/phones/" + file)

    for file in os.listdir(folders["slippers"]):
        imageId.append("products/slippers/" + file)

    for file in os.listdir(folders["shoes"]):
        imageId.append("products/shoes/" + file)

    for file in os.listdir(folders["watchs"]):
        imageId.append("products/watchs/" + file)


    selectimg = (random.sample(imageId, 12))
    img_group = []
    for i in range(0, len(selectimg),3):
            img_group.append(selectimg[i:i+3])
            
    return render_template("products.html",
                           searchimg = searchimg,
                           imageId = img_group,
                           offmenu = offmenu,
                           image = imagename,)


@app.route("/products/<path:image>")
def products(image):
    

    offmenu = True

    imagename = image
    imagedata = image.split("/")
    imgname = imagedata[2].split(".")
    phones = imagedata[1]
    hasSize = ""
    if phones in ["phones", "mainimg"]:
         hasSize = phones
    p_data = None
    for item in product:
        if str(imgname[0]) == str(item["p_id"]):
           p_data = item
           break

    imageId = []
    for file in os.listdir(folders["mens"]):
        imageId.append("products/mens_dresses/" + file)
    
    for file in os.listdir(folders["womens"]):
        imageId.append("products/women_dresses/" + file)

    for file in os.listdir(folders["phones"]):
        imageId.append("products/phones/" + file)

    for file in os.listdir(folders["slippers"]):
        imageId.append("products/slippers/" + file)

    for file in os.listdir(folders["shoes"]):
        imageId.append("products/shoes/" + file)

    for file in os.listdir(folders["watchs"]):
        imageId.append("products/watchs/" + file)

    if imagename == "products/img/":
        num = 15
    else:
        num = 9
       
       
    reviews = random.randrange(100, 281)
    
    selectimg = (random.sample(imageId, num))
    img_group = []
    for i in range(0, len(selectimg),3):
            img_group.append(selectimg[i:i+3])

    return render_template("products.html", 
                           imagedata2 = imagedata,
                           image = image,
                           imageId = img_group,
                           products = p_data,
                           hasSize = hasSize,
                           reviews = reviews,
                           offmenu = offmenu,)



@app.route("/register" , methods=["POST"])
def register():
     UI = session.get("userid")
     if UI:
          return jsonify({
                         "ok":True
                    })

     data = request.get_json()
             
     userName = data["username"]
     email = data["email"]
     password = data["password"]
     roll = "User"
    #  conform_password = request.form["conform-password"]
     if userName == "" or email == "" or password == "":
          return jsonify({
                        "ok":False
                    })
     else:
        excel_file = "Shopping_W_data.xlsx"
        sheet_name = "users_sheet"
        file = "Shopping_W_data.xlsx"
        

        wb = load_workbook(file)
        sheet1 = wb[sheet_name]
        last_userid = sheet1.cell(row=sheet1.max_row, column=1).value

        # Heading matrame unte
        if last_userid is None or last_userid == "":
            userid = "UI001"

        else:

            number = int(last_userid[2:])      # UI003 -> 3

            userid = f"UI{number+1:03d}"       # 4 -> UI004

        #  session["user"] = userName
        

        new_data = {
            "userid": [userid],
            "user" : [userName],
            "email" : [email],
            "password" : [password],
            "roll" : [roll],
        }
        session["userid"] = userid
        session["u_Name"] = userName

        new_df = pd.DataFrame(new_data)
        

        if os.path.exists(excel_file):
            book = load_workbook(excel_file)
            sheet = book[sheet_name]

            next_row = sheet.max_row
            with pd.ExcelWriter(excel_file,
                                engine="openpyxl",
                                mode="a",
                                if_sheet_exists="overlay") as writer:
                new_df.to_excel(
                        writer,
                        sheet_name=sheet_name,
                        index=False,
                        header=False,
                        startrow=next_row
                )
                print("user saved successfully")
                if data:
                        return jsonify({
                            "ok":True
                        })


@app.route("/login_page")
def login_page():
     userid = session.get("userid")

     if userid:
          return redirect("/")
     
     return render_template("login_page.html")

@app.route("/profile")
def profile():
    OdataF = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_history")
    OdataF.columns = OdataF.columns.str.strip().str.lower()
    OdataF = OdataF.dropna(how="all")
    UHproduct = OdataF.to_dict(orient="records")
    
    if not session.get("userid"):
         return render_template("login_page.html")
    userid = session.get("userid")
    df = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_sheet")
    df.columns = df.columns.str.strip().str.lower()
    df = df.dropna(how="all")
    userdata = df[df["userid"]== userid]

    userid = userdata.iloc[0]["userid"]
    username = userdata.iloc[0]["username"]
    email = userdata.iloc[0]["email"]

    totalOrders=""
    Pimgid =[]
    for item in UHproduct:
        if str(userid) == str(item["userid"]):
            Pimgid.append(item["p_id"])
            totalOrders = len(Pimgid)
            
    pr_Data={
        "username":username,
        "email":email,
        "totalOrders":totalOrders,
    }

    
    return render_template("profile.html", pr_Data=pr_Data)

@app.route("/check_login", methods=["POST"])
def check_login():
     df = pd.read_excel("Shopping_W_data.xlsx",
                        sheet_name="users_sheet")
     df.columns = df.columns.str.strip().str.lower()
     df = df.dropna(how="all")
      
     data = request.get_json()
     email = data["Lemail"]
     password = data["Lpassword"]

     

     if email in df["email"].values:
          userdata = df[df["email"]== email]
          excel_pwd = userdata.iloc[0]["password"]
          if userdata.empty:
           return jsonify({
               "success": False
          })
          elif password == excel_pwd:
                 session["userid"] = userdata.iloc[0]["userid"]
                 session["u_Name"] = userdata.iloc[0]["username"]
                 session["roll"] = userdata.iloc[0]["roll"]
                 return jsonify({
                     "success": True
                 })
          else: 
                 return jsonify({
                     "success": False
                 })
     else: 
        return jsonify({
            "success": False
                })
     
          



@app.route("/check-LR", methods=["POST"])
def check_email():

    data = request.get_json()
    email = data["email"]
    

    df = pd.read_excel("Shopping_W_data.xlsx",
                        sheet_name="users_sheet")
    df.columns = df.columns.str.strip().str.lower()
    df = df.dropna(how="all")
                        

    if email in df["email"].values:

        return jsonify({
            "message": "Email Already Exists"
        })

    else:

        return jsonify({
             
            "message": "Email Available"
        })



@app.route("/check_buydata", methods=["POST"])
def check_buydata():
    data = request.get_json()
    quantity = data["quantity"]
    size = data["Size"]
    imgpath = data["imgpath"]


    
    check_Q = quantity >= 1 and quantity <= 25
    check_S = size.lower() in ["s","m","l","xl","nosize"]
    check_imgp = bool(imgpath)
    
    if check_Q and check_S and check_imgp:
         
         session["buy_data"]={
              "quantity":quantity,
              "size" : size,
              "imgpath": imgpath
         }
         return jsonify({"ok":True})
    else:
         return jsonify({"ok": False})
    

@app.route("/buy_page")
def buy_page():
    buy_data = session.get("buy_data")

    if not buy_data:
         return redirect("/")

    imgpath = buy_data["imgpath"].split("/")

    imgID = imgpath[2].split(".")

    df = pd.read_excel("Shopping_W_data.xlsx")
    df.columns = df.columns.str.strip().str.lower()
    df = df.dropna(how="all")
    product = df.to_dict(orient="records")
    
    img_data = ""
    for item in product:
         if str(imgID[0]) == str(item["p_id"]):
              img_data = item
              break
         
    totat=int(img_data["p_price"]) * int(buy_data["quantity"])

    buy_data["totat_amount"] = totat


    date_time = datetime.now()
    O_date = date_time.strftime("%d/%m/%y")
    D_d = date_time.strftime("%d")
    D_date =(f"{3+int(D_d)}{date_time.strftime("/%m/%y")}")
    O_D_date = (f"{O_date}||{D_date}")

    

    
    
    return render_template("buy_page.html" 
                           ,buy_data = buy_data
                            ,img_data = img_data
                             ,O_date = O_date
                              ,D_date = D_date
                               ,O_D_date = O_D_date  )

@app.route("/write_excel", methods=["POST"])
def write_excel():
    #  users_history sheet
    OdataF = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_history")
    OdataF.columns = OdataF.columns.str.strip().str.lower()
    OdataF = OdataF.dropna(how="all")
    UHproduct = OdataF.to_dict(orient="records")

    
    data = request.get_json()

    imgpath = data["imgpath"].split("/")
    imgID = imgpath[2].split(".")

    quantity = data["quantity"]
    size = data["Size"]
    date = data["ODdate"]
    userID = session.get("userid")

    O_ID =[]
    for item in UHproduct:
        if str(userID) == str(item["userid"]):
            O_ID.append({
                 "OrderID" : item["orderid"]
            })

    if O_ID:
         last_order = O_ID[-1]["OrderID"]
         num = int(last_order.split("O")[-1])
         orderID = f"{userID}O{num+1}"
    else:
         orderID = (f"{userID}O1")


    excel_file = "Shopping_W_data.xlsx"
    sheet_name = "users_history"
    
    BUY_data = {
         "orderID": [orderID],
         "user_id": [userID],
         "Date": [date],
         "P_ID": [imgID[0]],
         "Qantity": [quantity],
         "Size": [size],
         "Status": "Ordered"
    }
    new_data= pd.DataFrame(BUY_data)

    if os.path.exists(excel_file):
         book = load_workbook(excel_file)
         sheet = book[sheet_name]
         
         next_row = sheet.max_row

         with pd.ExcelWriter(excel_file
                             ,engine="openpyxl"
                             ,mode="a"
                             ,if_sheet_exists="overlay") as writer:
              new_data.to_excel(
                   writer,
                   sheet_name=sheet_name,
                   index=False,
                   header=False,
                   startrow=next_row
              )
              session.pop("buy_data", None)
              print("buy_data save success")

    


    if data:
         return jsonify({"ok": True})


    return redirect("/")

@app.route("/logOut", methods=["POST"])
def logOut():
     data = request.get_json()
     logout = data["logout"]
     if logout:
          print("logOUT success")
          session.clear()
          return jsonify({"ok":True})
     return "logOUT success"

@app.route("/removeOP", methods=["POST"])
def removeOP():
     data = request.get_json()
     Rorderid = data["Rorderid"]

     XLname ="Shopping_W_data.xlsx"
     df = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_history")
     df.columns = df.columns.str.strip().str.lower()
     df = df.dropna(how="all")
     

     if Rorderid in df["orderid"].values:
        df.loc[df["orderid"] == Rorderid, "status"]= "Cancelled"
        with pd.ExcelWriter(
            XLname,
            engine="openpyxl",
            mode="a",
            if_sheet_exists="replace"
        ) as writer:
         df.to_excel(
            writer,
            sheet_name="users_history",
            index=False
        )
        print("Order Deleted successfully. ")
        return jsonify({"ok":True})
     else:
         print("Order ID not found. ")
         return jsonify({"ok":False})

@app.route("/Dashboard")
def Dashboard():
    if session.get("roll") == "Admin":
        userid = session.get("userid")
    else:
        return redirect("/")

    Odata = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_history")
    Odata.columns = Odata.columns.str.strip().str.lower()
    Odata = Odata.dropna(how="all")
    Odata = Odata.to_dict(orient="records")

    userdata = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_sheet")
    userdata.columns = userdata.columns.str.strip().str.lower()
    userdata = userdata.dropna(how="all")
    userdata = userdata.to_dict(orient="records")

    imgpath={}
    imgID = []
    for item in os.listdir(folders["products"]):
        path = os.path.join(folders["products"], item)
        img_lo = os.path.basename(path)
        for item1 in os.listdir(path):
            imgpath[item1.split(".")[0]]=f"{img_lo}/{item1}"
            imgID.append(item1.split(".")[0])

    
    for item in product:
        for item_id in imgID:
            if str(item["p_id"]) == str(item_id):
                item["img_path"]="products/" + imgpath[str(item_id)]

    data = {
        "total_O":len(Odata),
        "totalitem":len(product),
        "total_users":len(userdata),
        "users_D":userdata,
        "products":product,
        "Order_data":Odata,
        "imgpath":imgpath
    }

    
    return render_template("Dashboard.html",
                           data=data)

@app.route("/D_users_data/<Ddata>/<Ddata1>")
def D_users_data(Ddata,Ddata1):
    if session.get("roll") == "Admin":
        userid = session.get("userid")
    else:
        return redirect("/")

    if Ddata1 == "userdata":
        userdata = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_sheet")
        userdata.columns = userdata.columns.str.strip().str.lower()
        userdata = userdata.dropna(how="all")
        userdata = userdata.to_dict(orient="records")

        for item in userdata:
            if str(item["userid"]) == str(Ddata):
                    webdata={
                            "datatype":Ddata1,
                            "userid":Ddata,
                            "username":item["username"],
                            "email":item["email"],
                            "password":item["password"],
                            "roll":item["roll"],
                        }
                    break
    elif Ddata1 == "orderdata":
        Odata = pd.read_excel("Shopping_W_data.xlsx", sheet_name="users_history")
        Odata.columns = Odata.columns.str.strip().str.lower()
        Odata = Odata.dropna(how="all")
        Odata = Odata.to_dict(orient="records")

        
        for item in Odata:
            if str(item["orderid"]) == str(Ddata):
                 webdata={
                            "datatype":Ddata1,
                            "userid":item["userid"],
                            "orderid":Ddata,
                            "date":item["date"],
                            "quantity":item["quantity"],
                            "status":item["status"],
                        }
                 break

       
    else:
        print("product data")
    userID=Ddata
    
    return render_template("D_users_data.html",webdata=webdata)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug= True)
    

